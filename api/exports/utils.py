from __future__ import annotations

from io import BytesIO
import re
import zipfile

from reportlab.lib.pagesizes import LETTER
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.pdfgen import canvas
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt
import html

EXPORT_FONT_SIZE = 12
PDF_CJK_FONT = 'STSong-Light'
PDF_LATIN_FONT = 'Times-Roman'


def _escape_text(s: str) -> str:
    # Basic HTML escape to avoid accidental HTML rendering in downstream tools
    return html.escape(s, quote=False)


def proposal_json_to_markdown(proposal: dict) -> str:
    meta = proposal.get('meta', {})
    sections = proposal.get('sections', {}) or {}
    lines = []
    title = _escape_text(meta.get('title') or 'Proposal')
    lines.append(f'# {title}')
    for key, section in sections.items():
        title = _escape_text(section.get('title') or key)
        content = _escape_text(str(section.get('content') or ''))
        lines.append('')
        lines.append(f'## {title}')
        lines.append(content)
    return '\n'.join(lines)


def _normalize_pdf_for_checksum(data: bytes) -> bytes:
    # Remove/normalize non-deterministic parts: document ID and xref offset
    try:
        data = re.sub(rb'/ID\s*\[\s*<[^>]*>\s*<[^>]*>\s*\]', b'/ID [<000000><000000>]', data)
        data = re.sub(rb'startxref\s*\d+', b'startxref 0', data)
        data = re.sub(rb'/CreationDate\s*\(D:[^\)]+\)', b'/CreationDate (D:19700101000000Z)', data)
        data = re.sub(rb'/ModDate\s*\(D:[^\)]+\)', b'/ModDate (D:19700101000000Z)', data)
    except Exception:
        pass
    return data


def _register_pdf_fonts() -> tuple[str, str]:
    if PDF_CJK_FONT not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(UnicodeCIDFont(PDF_CJK_FONT))
    return PDF_CJK_FONT, PDF_LATIN_FONT


def _is_cjk(character: str) -> bool:
    codepoint = ord(character)
    return (
        0x2E80 <= codepoint <= 0x9FFF
        or 0xF900 <= codepoint <= 0xFAFF
        or 0xFF00 <= codepoint <= 0xFFEF
    )


def _wrap_pdf_line(text: str, max_width: float, cjk_font: str, latin_font: str, font_size: int = EXPORT_FONT_SIZE) -> list[str]:
    if not text:
        return ['']
    lines = []
    current = ''
    current_width = 0.0
    for character in text:
        font_name = cjk_font if _is_cjk(character) else latin_font
        width = pdfmetrics.stringWidth(character, font_name, font_size)
        if current and current_width + width > max_width:
            lines.append(current)
            current = character
            current_width = width
        else:
            current += character
            current_width += width
    lines.append(current)
    return lines


def _draw_mixed_font_line(c, x: float, y: float, line: str, cjk_font: str, latin_font: str, font_size: int = EXPORT_FONT_SIZE) -> None:
    text_object = c.beginText(x, y)
    active_font = None
    buffer = ''
    for character in line:
        font_name = cjk_font if _is_cjk(character) else latin_font
        if active_font is not None and font_name != active_font:
            text_object.setFont(active_font, font_size)
            text_object.textOut(buffer)
            buffer = ''
        active_font = font_name
        buffer += character
    if buffer:
        text_object.setFont(active_font or latin_font, font_size)
        text_object.textOut(buffer)
    c.drawText(text_object)


def _strip_inline_markdown(text: str) -> str:
    text = re.sub(r'!\[([^\]]*)\]\([^)]*\)', r'\1', text)
    text = re.sub(r'\[([^\]]+)\]\([^)]*\)', r'\1', text)
    return re.sub(r'(\*\*|__|~~|`)', '', text).strip()


def _markdown_pdf_blocks(markdown: str):
    for raw in str(markdown or '').splitlines():
        line = raw.rstrip()
        if not line.strip():
            yield 'spacer', ''
            continue
        heading = re.match(r'^(#{1,3})\s+(.+)$', line)
        if heading:
            yield f'heading_{len(heading.group(1))}', _strip_inline_markdown(heading.group(2))
            continue
        bullet = re.match(r'^\s*[-*+]\s+(.+)$', line)
        if bullet:
            yield 'list', '• ' + _strip_inline_markdown(bullet.group(1))
            continue
        numbered = re.match(r'^\s*(\d+[.)])\s+(.+)$', line)
        if numbered:
            yield 'list', f'{numbered.group(1)} {_strip_inline_markdown(numbered.group(2))}'
            continue
        quote = re.match(r'^\s*>\s?(.+)$', line)
        if quote:
            yield 'quote', _strip_inline_markdown(quote.group(1))
            continue
        yield 'paragraph', _strip_inline_markdown(line)


def render_pdf_from_text(text: str) -> tuple[bytes, str]:
    buffer = BytesIO()
    c = canvas.Canvas(buffer, pagesize=LETTER)
    cjk_font, latin_font = _register_pdf_fonts()
    # Deterministic metadata (guarded: ReportLab internals may differ across versions)
    try:  # Accessing protected internals for deterministic metadata; ignore type
        info = c._doc.info  # type: ignore[attr-defined]
        info.title = 'FoundationPal Export'
        info.author = 'FoundationPal'
        info.creator = 'FoundationPal'
        info.producer = 'ReportLab'
        info.creationDate = 'D:19700101000000Z'
        info.modDate = 'D:19700101000000Z'
    except Exception:  # pragma: no cover - defensive
        pass
    width, height = LETTER
    y = height - 72
    styles = {
        'heading_1': (18, 25, 72),
        'heading_2': (16, 23, 72),
        'heading_3': (14, 21, 72),
        'paragraph': (EXPORT_FONT_SIZE, 18, 72),
        'list': (EXPORT_FONT_SIZE, 18, 88),
        'quote': (EXPORT_FONT_SIZE, 18, 84),
    }
    for kind, line in _markdown_pdf_blocks(text):
        if kind == 'spacer':
            y -= 9
            continue
        font_size, line_height, x = styles[kind]
        for wrapped_line in _wrap_pdf_line(line, width - x - 72, cjk_font, latin_font, font_size):
            if y < 72 + line_height:
                c.showPage()
                y = height - 72
            _draw_mixed_font_line(c, x, y, wrapped_line, cjk_font, latin_font, font_size)
            y -= line_height
        if kind.startswith('heading'):
            y -= 4
    c.save()
    data = buffer.getvalue()
    # Post-process raw PDF to enforce epoch Creation/Mod dates so raw bytes deterministic
    try:
        data = re.sub(rb'/CreationDate\s*\(D:[^\)]+\)', b'/CreationDate (D:19700101000000Z)', data)
        data = re.sub(rb'/ModDate\s*\(D:[^\)]+\)', b'/ModDate (D:19700101000000Z)', data)
    except Exception:  # pragma: no cover - defensive
        pass
    normalized = _normalize_pdf_for_checksum(data)
    # Defer to common checksum utility for consistency
    try:
        from app.common.files import compute_checksum  # local import to avoid early app loading side-effects

        checksum = compute_checksum(normalized).hex
    except Exception:
        # Fallback to direct hashlib if utility unavailable (defensive during migrations)
        import hashlib as _hl

        checksum = _hl.sha256(normalized).hexdigest()
    return data, checksum


def render_docx_from_markdown(md: str) -> tuple[bytes, str]:
    """
    Minimal Markdown -> DOCX renderer.
    Supports:
      - # H1 and ## H2 headings
      - Paragraphs
      - Blank lines for spacing
    This keeps the canonical markdown as the source of truth and does a light render.
    """
    doc = Document()
    def apply_run_font(run) -> None:
        run.font.name = 'Times New Roman'
        run.font.size = Pt(EXPORT_FONT_SIZE)
        run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '宋体')

    # Apply the same Chinese and English fonts to body text and headings.
    try:  # Styles may not always be mutable
        for style_name in ('Normal', 'Title', 'Heading 1', 'Heading 2'):
            style = doc.styles[style_name]  # type: ignore[index]
            style.font.name = 'Times New Roman'  # type: ignore[attr-defined]
            style.font.size = Pt(EXPORT_FONT_SIZE)  # type: ignore[attr-defined]
            style._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '宋体')
    except Exception:  # pragma: no cover - defensive
        pass

    for raw in md.splitlines():
        line = raw.rstrip()
        if not line:
            doc.add_paragraph('')
            continue
        if line.startswith('# '):
            p = doc.add_heading(line[2:].strip(), level=1)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for run in p.runs:
                apply_run_font(run)
            continue
        if line.startswith('## '):
            p = doc.add_heading(line[3:].strip(), level=2)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for run in p.runs:
                apply_run_font(run)
            continue
        paragraph = doc.add_paragraph(line)
        for run in paragraph.runs:
            apply_run_font(run)

    # Deterministic core properties
    try:
        core = doc.core_properties
        core.title = 'FoundationPal Export'
        core.author = 'FoundationPal'
        from datetime import datetime, timezone

        epoch = datetime(1970, 1, 1, tzinfo=timezone.utc)
        core.created = epoch
        core.modified = epoch
        core.last_printed = epoch
    except Exception:
        pass

    bio = BytesIO()
    doc.save(bio)
    raw = bio.getvalue()

    # Normalize DOCX zip for determinism: fixed timestamps and sorted entries
    def _normalize_docx_zip(data: bytes) -> bytes:
        src = BytesIO(data)
        out_bio = BytesIO()
        with zipfile.ZipFile(src, 'r') as zin, zipfile.ZipFile(out_bio, 'w', compression=zipfile.ZIP_DEFLATED) as zout:
            for name in sorted(zin.namelist()):
                content = zin.read(name)
                zi = zipfile.ZipInfo(filename=name, date_time=(1980, 1, 1, 0, 0, 0))
                zi.compress_type = zipfile.ZIP_DEFLATED
                # Ensure consistent permissions/external attrs
                zi.external_attr = 0o600 << 16
                zout.writestr(zi, content)
        return out_bio.getvalue()

    data = _normalize_docx_zip(raw)
    try:
        from app.common.files import compute_checksum

        checksum = compute_checksum(data).hex
    except Exception:
        import hashlib as _hl

        checksum = _hl.sha256(data).hexdigest()
    return data, checksum
