"""NSFC基金申请Agent - Gradio Demo"""
import json, urllib.request, os, tempfile, re

API = "http://127.0.0.1:8000/api/ai"
SECTION_IDS = ["lixiangyiju", "yanjiuneirong", "yanjiufangan", "yanjiujichu"]
SECTION_TITLES = ["(一)立项依据", "(二)研究内容", "(三)研究方案", "(四)研究基础"]
BOXES_PER_SECTION = 20
LOADING_TEXT = "正在处理中，请耐心等待..."

def call_api(endpoint, data, timeout=180):
    d = json.dumps(data).encode()
    req = urllib.request.Request(f"{API}/{endpoint}", data=d, headers={"Content-Type": "application/json"})
    return json.loads(urllib.request.urlopen(req, timeout=timeout).read())

def save_as_md(content):
    path = os.path.join(tempfile.gettempdir(), "nsfc_draft.md")
    with open(path, "w", encoding="utf-8") as f:
        f.write(content)
    return path

def save_as_txt(content):
    path = os.path.join(tempfile.gettempdir(), "nsfc_draft.txt")
    text = re.sub(r"#+\s*", "", content)
    text = re.sub(r"\*\*|\*", "", text)
    with open(path, "w", encoding="utf-8") as f:
        f.write(text)
    return path

def save_as_docx(content):
    try:
        from docx import Document
        path = os.path.join(tempfile.gettempdir(), "nsfc_draft.docx")
        doc = Document()
        for line in content.split("\n"):
            if line.startswith("## "):
                doc.add_heading(line[3:], level=2)
            elif line.startswith("### "):
                doc.add_heading(line[4:], level=3)
            elif line.strip():
                doc.add_paragraph(line)
        doc.save(path)
        return path
    except ImportError:
        return ""

def extract_file(file_path):
    if not file_path:
        return "", ""
    ext = os.path.splitext(file_path)[1].lower()
    try:
        if ext in (".txt", ".md"):
            with open(file_path, "r", encoding="utf-8") as f:
                text = f.read()
        elif ext == ".docx":
            from docx import Document
            doc = Document(file_path)
            text = "\n".join([p.text for p in doc.paragraphs if p.text.strip()])
        elif ext == ".pdf":
            import pdfplumber
            with pdfplumber.open(file_path) as pdf:
                text = "\n".join([page.extract_text() or "" for page in pdf.pages])
        else:
            text = ""
    except Exception as e:
        text = f"[文件读取失败: {e}]"
    return text, text

def step1_plan(research_topic):
    empty_updates = [gr.update()] * (4 * BOXES_PER_SECTION)
    if not research_topic.strip():
        return ("请输入研究方向", "", "", *empty_updates, "")
    result = call_api("plan", {"grant_url": "", "text_spec": research_topic})
    sections = result.get("plan", {}).get("sections", [])
    questions_text = "\n".join(
        f"## {s['title']}\n" + "\n".join(f"- {q}" for q in s.get("questions", []))
        for s in sections
    )
    updates = []
    for sec_id in SECTION_IDS:
        sec = next((s for s in sections if s["id"] == sec_id), None)
        qs = sec["questions"] if sec else []
        for i in range(BOXES_PER_SECTION):
            if i < len(qs):
                updates.append(gr.update(visible=True, label=qs[i], value=""))
            else:
                updates.append(gr.update(visible=False, label="", value=""))
    return (
        f"已生成{len(sections)}个章节，切换到Tab2回答问题",
        json.dumps(sections, ensure_ascii=False, indent=2),
        questions_text,
        *updates,
        "",
    )

def step2_write(sections_json_str, *all_answers):
    try:
        sections = json.loads(sections_json_str) if sections_json_str.strip() else []
    except:
        return "规划数据解析失败", "", "", "", ""
    results = []
    for idx, sec_id in enumerate(SECTION_IDS):
        sec = next((s for s in sections if s["id"] == sec_id), None)
        if not sec:
            continue
        qs = sec["questions"]
        answers = {}
        for i in range(min(len(qs), BOXES_PER_SECTION)):
            val = all_answers[idx * BOXES_PER_SECTION + i]
            if val and val.strip():
                answers[qs[i]] = val.strip()
        if not answers:
            continue
        try:
            r = call_api("write", {"section_id": sec_id, "answers": answers}, timeout=180)
            draft = r.get("draft_text", "")
            results.append(f"## {sec['title']}\n\n{draft}")
        except Exception as e:
            results.append(f"## {sec['title']}\n\n[错误: {e}]")
    if not results:
        return "请先在Tab2中填写至少一个章节的回答", "", "", "", ""
    full = "\n\n---\n\n".join(results)
    return full, full, save_as_md(full), save_as_txt(full), save_as_docx(full) or ""

def step3_revise(draft_text, change_request):
    if not draft_text.strip() or not change_request.strip():
        return "请填写原文和修改要求", "", "", "", ""
    result = call_api("revise", {
        "base_text": draft_text, "change_request": change_request, "section_id": "revision",
    }, timeout=180)
    revised = result.get("draft_text", "")
    return revised, revised, save_as_md(revised), save_as_txt(revised), save_as_docx(revised) or ""

def step4_format(full_text):
    if not full_text.strip():
        return "请先完成撰写", "", "", "", ""
    result = call_api("format", {"full_text": full_text}, timeout=180)
    formatted = result.get("formatted_text", "")
    return formatted, formatted, save_as_md(formatted), save_as_txt(formatted), save_as_docx(formatted) or ""


if __name__ == "__main__":
    import gradio as gr
    from functools import partial

    with gr.Blocks(title="NSFC基金申请Agent") as demo:
        gr.Markdown("# NSFC 基金申请书撰写助手")

        draft_state = gr.State("")
        revised_state = gr.State("")

        # ======== Tab 1 ========
        with gr.Tab("1. 规划章节"):
            topic = gr.Textbox(label="研究方向", lines=2, placeholder="例如:基于深度学习的自动调制识别方法研究")
            with gr.Row():
                btn_plan = gr.Button("生成章节规划", variant="primary")
            plan_status = gr.Markdown("")
            with gr.Row():
                sec_json = gr.Code(label="章节结构JSON", language="json", lines=6)
                questions = gr.Markdown()

        # ======== Tab 2 ========
        with gr.Tab("2. 撰写内容"):
            all_boxes = []
            for sec_title in SECTION_TITLES:
                with gr.Accordion(sec_title, open=True):
                    for i in range(BOXES_PER_SECTION):
                        box = gr.Textbox(label=f"Q{i+1}", visible=False, lines=2)
                        all_boxes.append(box)
            with gr.Row():
                btn_load = gr.Button("加载规划结果", variant="secondary")
                btn_write = gr.Button("生成申请书草稿", variant="primary")
            write_status = gr.Markdown("")
            draft_md = gr.Markdown()
            with gr.Row():
                dl_md = gr.DownloadButton("下载 .md", visible=False)
                dl_txt = gr.DownloadButton("下载 .txt", visible=False)
                dl_docx = gr.DownloadButton("下载 .docx", visible=False)

        # ======== Tab 3 ========
        with gr.Tab("3. 修改润色"):
            gr.Markdown("加载已撰写内容或上传附件，输入修改要求进行修订。")
            with gr.Row():
                btn_load_draft = gr.Button("加载撰写内容", variant="secondary")
                upload_file = gr.File(label="上传附件", file_types=[".txt", ".md", ".docx", ".pdf"])
            draft_input = gr.Textbox(label="原文(可编辑)", lines=12)
            change_req = gr.Textbox(label="修改要求", lines=3, placeholder="例如:加强创新点论述，补充近三年文献")
            btn_revise = gr.Button("修订", variant="primary")
            revise_status = gr.Markdown("")
            revised_md = gr.Markdown()
            with gr.Row():
                r_dl_md = gr.DownloadButton("下载 .md", visible=False)
                r_dl_txt = gr.DownloadButton("下载 .txt", visible=False)
                r_dl_docx = gr.DownloadButton("下载 .docx", visible=False)

        # ======== Tab 4 ========
        with gr.Tab("4. 排版输出"):
            with gr.Row():
                btn_load_revised = gr.Button("加载修订后的结果", variant="secondary")
            full_input = gr.Textbox(label="全文(可编辑)", lines=18)
            btn_format = gr.Button("格式化输出", variant="primary")
            fmt_status = gr.Markdown("")
            formatted_md = gr.Markdown()
            with gr.Row():
                f_dl_md = gr.DownloadButton("下载 .md", visible=False)
                f_dl_txt = gr.DownloadButton("下载 .txt", visible=False)
                f_dl_docx = gr.DownloadButton("下载 .docx", visible=False)

        # ======== Events ========

        btn_plan.click(
            fn=lambda: LOADING_TEXT, outputs=[plan_status]
        ).then(
            fn=step1_plan, inputs=[topic],
            outputs=[plan_status, sec_json, questions] + all_boxes + [draft_state],
        )

        btn_load.click(
            fn=lambda: LOADING_TEXT, outputs=[plan_status]
        ).then(
            fn=step1_plan, inputs=[topic],
            outputs=[plan_status, sec_json, questions] + all_boxes + [draft_state],
        )

        btn_write.click(
            fn=lambda: LOADING_TEXT, outputs=[write_status]
        ).then(
            fn=step2_write, inputs=[sec_json] + all_boxes,
            outputs=[draft_md, draft_state, dl_md, dl_txt, dl_docx],
        ).then(
            fn=lambda: "", outputs=[write_status],
        )

        btn_load_draft.click(
            fn=lambda d: d, inputs=[draft_state], outputs=[draft_input],
        )

        upload_file.upload(
            fn=extract_file, inputs=[upload_file], outputs=[draft_input, revised_state],
        )

        btn_revise.click(
            fn=lambda: LOADING_TEXT, outputs=[revise_status]
        ).then(
            fn=step3_revise, inputs=[draft_input, change_req],
            outputs=[revised_md, revised_state, r_dl_md, r_dl_txt, r_dl_docx],
        ).then(
            fn=lambda: "", outputs=[revise_status],
        )

        btn_load_revised.click(
            fn=lambda r: r, inputs=[revised_state], outputs=[full_input],
        )

        btn_format.click(
            fn=lambda: LOADING_TEXT, outputs=[fmt_status]
        ).then(
            fn=step4_format, inputs=[full_input],
            outputs=[formatted_md, revised_state, f_dl_md, f_dl_txt, f_dl_docx],
        ).then(
            fn=lambda: "", outputs=[fmt_status],
        )

    demo.launch(server_name="127.0.0.1", server_port=7860, share=False)
