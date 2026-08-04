import io
import tempfile

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from docx import Document
from docx.enum.text import WD_BREAK
from reportlab.pdfgen import canvas
from rest_framework.test import APIClient

from ai.models import AIResource, EvidenceFact, EvidenceUsage, UserEvidence
from ai.project_materials import create_project_material
from orgs.models import Organization, OrgUser
from proposals.models import Proposal, ProposalSection


def three_page_pdf():
    output = io.BytesIO()
    document = canvas.Canvas(output)
    for page_number in range(1, 4):
        document.drawString(72, 720, f'uploaded evidence page {page_number}')
        if page_number < 3:
            document.showPage()
    document.save()
    return output.getvalue()


def three_page_docx():
    output = io.BytesIO()
    document = Document()
    document.add_paragraph('uploaded docx page 1')
    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    document.add_paragraph('uploaded docx page 2')
    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    document.add_paragraph('uploaded docx page 3')
    document.save(output)
    return output.getvalue()


@override_settings(DEBUG=True, AI_PROVIDER='stub', AI_ASYNC=False)
class ProjectSetupTests(TestCase):
    def setUp(self):
        user_model = get_user_model()
        self.user = user_model.objects.create_user(username='setup-user', password='p')
        self.org = Organization.objects.create(name='Setup Org', admin=self.user)
        OrgUser.objects.create(org=self.org, user=self.user, role='admin')
        self.proposal = Proposal.objects.create(author=self.user, org=self.org, content={'meta': {'title': '测试项目'}})
        self.api = APIClient()
        self.api.force_authenticate(self.user)
        self.headers = {'HTTP_X_ORG_ID': str(self.org.id)}

    def test_setup_creates_project_scoped_guideline_and_confirmed_brief_evidence(self):
        response = self.api.post(
            f'/api/ai/proposals/{self.proposal.id}/setup',
            {
                'research_direction': '面向低信噪比通信信号的自动调制识别方法。',
                'core_problem': '如何兼顾少样本条件下的识别准确性与跨场景泛化能力。',
                'guideline_text': '申请书应当说明研究目标、研究内容和技术路线。',
            },
            format='json',
            **self.headers,
        )

        self.assertEqual(response.status_code, 200, response.content)
        self.assertTrue(response.json()['ready'])
        self.assertEqual(response.json()['guideline_text'], '申请书应当说明研究目标、研究内容和技术路线。')
        self.assertEqual(AIResource.objects.filter(proposal_id=self.proposal.id, knowledge_domain='grant_rule').count(), 1)
        self.assertGreater(AIResource.objects.filter(proposal_id=self.proposal.id, knowledge_domain='user_evidence').count(), 0)
        self.assertTrue(UserEvidence.objects.filter(proposal=self.proposal).exists())
        self.assertTrue(EvidenceFact.objects.filter(user_evidence__proposal=self.proposal, verification_status='user_confirmed').exists())
        status = self.api.get(f'/api/ai/proposals/{self.proposal.id}/setup', **self.headers)
        self.assertEqual(status.status_code, 200, status.content)
        self.assertEqual(status.json()['guideline_text'], '申请书应当说明研究目标、研究内容和技术路线。')

    def test_same_material_is_scoped_to_each_project(self):
        second = Proposal.objects.create(author=self.user, org=self.org, content={'meta': {'title': '第二项目'}})
        payload = {
            'research_direction': '同一研究方向。',
            'guideline_text': '申请书应当说明研究目标。',
        }
        first = self.api.post(f'/api/ai/proposals/{self.proposal.id}/setup', payload, format='json', **self.headers)
        second_response = self.api.post(f'/api/ai/proposals/{second.id}/setup', payload, format='json', **self.headers)

        self.assertEqual(first.status_code, 200, first.content)
        self.assertEqual(second_response.status_code, 200, second_response.content)
        resources = AIResource.objects.filter(knowledge_domain='grant_rule').order_by('proposal_id')
        self.assertEqual(list(resources.values_list('proposal_id', flat=True)), [self.proposal.id, second.id])

    def test_resaving_setup_replaces_its_previous_pasted_materials(self):
        first = self.api.post(
            f'/api/ai/proposals/{self.proposal.id}/setup',
            {'research_direction': '第一版研究方向。', 'guideline_text': '第一版基金指南要求。'},
            format='json',
            **self.headers,
        )
        second = self.api.post(
            f'/api/ai/proposals/{self.proposal.id}/setup',
            {'research_direction': '第二版研究方向。', 'guideline_text': '第二版基金指南要求。'},
            format='json',
            **self.headers,
        )

        self.assertEqual(first.status_code, 200, first.content)
        self.assertEqual(second.status_code, 200, second.content)
        active = AIResource.objects.filter(proposal_id=self.proposal.id, is_deleted=False)
        self.assertEqual(active.filter(knowledge_domain='grant_rule').count(), 1)
        self.assertEqual(active.filter(knowledge_domain='user_evidence').count(), 1)
        self.assertEqual(second.json()['research_direction'], '第二版研究方向。')

    def test_uploaded_text_material_is_ingested_for_its_project(self):
        with tempfile.TemporaryDirectory() as media_root, override_settings(MEDIA_ROOT=media_root):
            response = self.api.post(
                '/api/files',
                {
                    'file': SimpleUploadedFile('基础材料.txt', '已有实验平台和采集数据。'.encode('utf-8'), content_type='text/plain'),
                    'proposal_id': str(self.proposal.id),
                    'material_kind': 'evidence',
                },
                format='multipart',
                **self.headers,
            )

        self.assertEqual(response.status_code, 200, response.content)
        self.assertEqual(response.json()['material_kind'], 'evidence')
        resource_id = response.json()['resource_id']
        self.assertTrue(UserEvidence.objects.filter(resource_id=resource_id, proposal=self.proposal).exists())

    def test_uploaded_material_keeps_its_real_page_while_project_setup_evidence_hides_pages(self):
        with tempfile.TemporaryDirectory() as media_root, override_settings(MEDIA_ROOT=media_root):
            response = self.api.post(
                '/api/files',
                {
                    'file': SimpleUploadedFile('supplement.pdf', three_page_pdf(), content_type='application/pdf'),
                    'proposal_id': str(self.proposal.id),
                    'material_kind': 'evidence',
                },
                format='multipart',
                **self.headers,
            )

            self.assertEqual(response.status_code, 200, response.content)
            resource = AIResource.objects.get(id=response.json()['resource_id'])
            third_page = resource.chunks.get(page_start=3)
            self.assertEqual(third_page.page_end, 3)
            section = ProposalSection.objects.create(proposal=self.proposal, key='evidence-page')
            EvidenceUsage.objects.create(
                proposal_section=section,
                chunk=third_page,
                role='writer',
                snapshot_text=third_page.text,
                document_name_snapshot='supplement.pdf',
                page_start_snapshot=1,
                page_end_snapshot=1,
            )
            evidence = self.api.get(f'/api/ai/sections/{section.id}/evidence', **self.headers).json()['evidence'][0]

        self.assertTrue(evidence['is_uploaded_material'])
        self.assertEqual((evidence['page_start'], evidence['page_end']), (3, 3))

        setup = self.api.post(
            f'/api/ai/proposals/{self.proposal.id}/setup',
            {'research_direction': '研究方向。', 'guideline_text': '指南要求。'},
            format='json',
            **self.headers,
        )
        self.assertEqual(setup.status_code, 200, setup.content)
        brief = AIResource.objects.filter(
            proposal_id=self.proposal.id,
            knowledge_domain='user_evidence',
            original_filename='',
        ).latest('id').chunks.first()
        EvidenceUsage.objects.create(
            proposal_section=section,
            chunk=brief,
            role='writer',
            snapshot_text=brief.text,
            document_name_snapshot='项目研究方向与核心问题',
        )
        payload = self.api.get(f'/api/ai/sections/{section.id}/evidence', **self.headers).json()['evidence']
        setup_evidence = next(item for item in payload if item['document_name'] == '项目研究方向与核心问题')
        self.assertFalse(setup_evidence['is_uploaded_material'])
        self.assertIsNone(setup_evidence['page_start'])
        self.assertIsNone(setup_evidence['page_end'])

    def test_uploaded_docx_uses_explicit_page_breaks_for_evidence_pages(self):
        with tempfile.TemporaryDirectory() as media_root, override_settings(MEDIA_ROOT=media_root):
            response = self.api.post(
                '/api/files',
                {
                    'file': SimpleUploadedFile(
                        'supplement.docx',
                        three_page_docx(),
                        content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
                    ),
                    'proposal_id': str(self.proposal.id),
                    'material_kind': 'evidence',
                },
                format='multipart',
                **self.headers,
            )

            self.assertEqual(response.status_code, 200, response.content)
            resource = AIResource.objects.get(id=response.json()['resource_id'])
            self.assertTrue(resource.chunks.filter(page_start=3, page_end=3).exists())

    def test_status_lists_only_this_projects_uploaded_materials(self):
        other = Proposal.objects.create(author=self.user, org=self.org, content={'meta': {'title': 'Other'}})
        create_project_material(
            proposal=self.proposal,
            owner=self.user,
            material_kind='guideline',
            title='guide-a.txt',
            original_filename='guide-a.txt',
            text='First project requirement.',
        )
        create_project_material(
            proposal=self.proposal,
            owner=self.user,
            material_kind='evidence',
            title='evidence-a.txt',
            original_filename='evidence-a.txt',
            text='First project evidence.',
        )
        create_project_material(
            proposal=other,
            owner=self.user,
            material_kind='guideline',
            title='guide-b.txt',
            original_filename='guide-b.txt',
            text='Other project requirement.',
        )

        response = self.api.get(f'/api/ai/proposals/{self.proposal.id}/setup', **self.headers)

        self.assertEqual(response.status_code, 200, response.content)
        payload = response.json()
        self.assertEqual(payload['guideline_count'], 1)
        self.assertEqual(payload['evidence_count'], 1)
        self.assertEqual([row['name'] for row in payload['guideline_files']], ['guide-a.txt'])
        self.assertEqual([row['name'] for row in payload['evidence_files']], ['evidence-a.txt'])

    def test_long_project_guideline_is_split_into_traceable_chunks(self):
        resource = create_project_material(
            proposal=self.proposal,
            owner=self.user,
            material_kind='guideline',
            title='long-guide.txt',
            original_filename='long-guide.txt',
            text='Requirement sentence. ' * 220,
        )

        chunks = list(resource.chunks.order_by('chunk_index'))
        self.assertGreater(len(chunks), 1)
        self.assertTrue(all(len(chunk.text) <= 1000 for chunk in chunks))
        self.assertEqual(resource.parser_version, 'project-material-v2')
